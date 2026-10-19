"""Conversions AikoRoue pour Windows et Linux.

Tout tourne en local : Pillow (images), pypdfium2 + pypdf (PDF),
ffmpeg fourni par imageio-ffmpeg (vidéo/audio), Qt (documents → PDF).
"""
from __future__ import annotations

import io
import os
import subprocess
import sys
import zipfile
from pathlib import Path

from PIL import Image, ImageOps

try:  # lecture/écriture HEIC
    import pillow_heif

    pillow_heif.register_heif_opener()
    HAS_HEIF = True
except Exception:  # pragma: no cover
    HAS_HEIF = False


class Oops(Exception):
    """Erreur affichée telle quelle à l'utilisateur."""


# ---------------------------------------------------------------- sorties

def out_file(src: Path, ext: str, suffix: str = "") -> Path:
    """« photo.png » → « photo.jpg » à côté, sans jamais écraser (« photo 2.jpg »)."""
    return unique(src.parent, src.stem + suffix, ext)


def unique(folder: Path, base: str, ext: str) -> Path:
    def make(b: str) -> Path:
        return folder / (f"{b}.{ext}" if ext else b)

    path, n = make(base), 2
    while path.exists():
        path, n = make(f"{base} {n}"), n + 1
    return path


def out_folder(src: Path, suffix: str) -> Path:
    folder = unique(src.parent, src.stem + suffix, "")
    folder.mkdir(parents=True)
    return folder


def size(p: Path) -> int:
    try:
        return p.stat().st_size
    except OSError:
        return 0


def discard(p: Path) -> None:
    try:
        p.unlink()
    except OSError:
        pass


# ---------------------------------------------------------------- ffmpeg

def ffmpeg_exe() -> str:
    try:
        import imageio_ffmpeg

        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        return "ffmpeg"


def run_ffmpeg(args: list[str], out: Path) -> None:
    cmd = [ffmpeg_exe(), "-hide_banner", "-loglevel", "error", "-y", *args, str(out)]
    flags = 0x08000000 if sys.platform == "win32" else 0  # CREATE_NO_WINDOW
    proc = subprocess.run(cmd, capture_output=True, text=True, creationflags=flags)
    if proc.returncode != 0:
        discard(out)
        lines = [l for l in (proc.stderr or "").strip().splitlines() if l.strip()]
        raise Oops(lines[-1] if lines else "ffmpeg a échoué")


# ---------------------------------------------------------------- images

IMAGE_SAVE = {  # ext → (format Pillow, options)
    "jpg": ("JPEG", {"quality": 92, "optimize": True}),
    "png": ("PNG", {"optimize": True}),
    "webp": ("WEBP", {"quality": 90, "method": 5}),
    "heic": ("HEIF", {"quality": 90}),
    "gif": ("GIF", {}),
    "tiff": ("TIFF", {"compression": "tiff_lzw"}),
    "bmp": ("BMP", {}),
    "ico": ("ICO", {"sizes": [(16, 16), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)]}),
}


def norm_ext(ext: str) -> str:
    ext = ext.lower().lstrip(".")
    return {"jpeg": "jpg", "jpe": "jpg", "tif": "tiff", "htm": "html", "aif": "aiff",
            "heif": "heic", "markdown": "md"}.get(ext, ext)


def load_image(path: Path) -> tuple[Image.Image, bytes | None]:
    try:
        im = Image.open(path)
        im.load()
    except Exception:
        raise Oops(f"Impossible de lire {path.name}")
    exif = None
    try:
        e = im.getexif()
        if e:
            e[0x0112] = 1  # l'orientation est appliquée ci-dessous
            exif = e.tobytes()
    except Exception:
        exif = None
    im = ImageOps.exif_transpose(im)
    return im, exif


def flatten(im: Image.Image) -> Image.Image:
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, (255, 255, 255))
        bg.paste(im, mask=im.getchannel("A"))
        return bg
    return im.convert("RGB") if im.mode not in ("RGB", "L") else im


def save_image(im: Image.Image, out: Path, ext: str, exif: bytes | None = None, quality: int | None = None) -> None:
    fmt, opts = IMAGE_SAVE[ext]
    opts = dict(opts)
    if quality is not None and "quality" in opts:
        opts["quality"] = quality
    if ext in ("jpg", "bmp"):
        im = flatten(im)
    elif ext in ("png", "webp", "tiff", "heic", "ico") and im.mode not in ("RGB", "RGBA", "L", "LA"):
        im = im.convert("RGBA")
    if exif and ext in ("jpg", "png", "webp", "heic"):
        opts["exif"] = exif
    try:
        im.save(out, fmt, **opts)
    except Exception as e:
        discard(out)
        raise Oops(f"Échec de l'écriture de {out.name} ({e})")


def same_image_ext(path: Path) -> str:
    ext = norm_ext(path.suffix)
    if ext in IMAGE_SAVE and ext != "ico" and (ext != "heic" or HAS_HEIF):
        return ext
    return "png"


def image_convert(paths: list[Path], ext: str) -> list[Path]:
    outs = []
    for p in paths:
        im, exif = load_image(p)
        out = out_file(p, ext)
        save_image(im, out, ext, exif)
        outs.append(out)
    return outs


def image_to_pdf(paths: list[Path]) -> list[Path]:
    pages = [flatten(load_image(p)[0]) for p in paths]
    suffix = f" (+{len(paths) - 1})" if len(paths) > 1 else ""
    out = out_file(paths[0], "pdf", suffix)
    pages[0].save(out, "PDF", save_all=True, append_images=pages[1:], resolution=150)
    return [out]


def image_compress(paths: list[Path]) -> list[Path]:
    outs = []
    for p in paths:
        im, exif = load_image(p)
        alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
        im.thumbnail((1920, 1920) if alpha else (2560, 2560), Image.LANCZOS)
        ext = "png" if alpha else "jpg"
        out = out_file(p, ext, " (compressé)")
        save_image(im, out, ext, exif, quality=65)
        if size(out) >= size(p):
            discard(out)
            continue
        outs.append(out)
    if not outs:
        raise Oops("Déjà optimisé : impossible de réduire davantage")
    return outs


def image_resize(paths: list[Path], scale: float = 0.5) -> list[Path]:
    outs = []
    for p in paths:
        im, exif = load_image(p)
        w, h = im.size
        im = im.resize((max(1, int(w * scale)), max(1, int(h * scale))), Image.LANCZOS)
        ext = same_image_ext(p)
        out = out_file(p, ext, " (réduit)")
        save_image(im, out, ext, exif)
        outs.append(out)
    return outs


def image_rotate(paths: list[Path]) -> list[Path]:
    outs = []
    for p in paths:
        im, exif = load_image(p)
        im = im.rotate(-90, expand=True)
        ext = same_image_ext(p)
        out = out_file(p, ext, " (pivoté)")
        save_image(im, out, ext, exif)
        outs.append(out)
    return outs


def image_strip(paths: list[Path]) -> list[Path]:
    outs = []
    for p in paths:
        im, _ = load_image(p)
        clean = im.copy()
        transparency = im.info.get("transparency")
        clean.info = {} if transparency is None else {"transparency": transparency}
        ext = same_image_ext(p)
        out = out_file(p, ext, " (sans métadonnées)")
        save_image(clean, out, ext, None)
        outs.append(out)
    return outs


# ---------------------------------------------------------------- PDF

def _pdfium(path: Path):
    import pypdfium2 as pdfium

    try:
        return pdfium.PdfDocument(str(path))
    except Exception:
        raise Oops(f"Impossible d'ouvrir {path.name} (protégé ?)")


def pdf_to_images(paths: list[Path], ext: str) -> list[Path]:
    outs = []
    for p in paths:
        doc = _pdfium(p)
        n = len(doc)
        folder = out_folder(p, f" ({ext.upper()})") if n > 1 else None
        for i in range(n):
            img = doc[i].render(scale=2).to_pil()
            out = unique(folder, f"Page {i + 1}", ext) if folder else out_file(p, ext)
            save_image(img, out, ext)
            if not folder:
                outs.append(out)
        if folder:
            outs.append(folder)
        doc.close()
    return outs


def _pdf_pages_text(path: Path) -> list[str]:
    doc = _pdfium(path)
    texts = []
    for i in range(len(doc)):
        tp = doc[i].get_textpage()
        texts.append(tp.get_text_range())
        tp.close()
    doc.close()
    if not "".join(texts).strip():
        raise Oops(f"Aucun texte dans {path.name} (PDF scanné ?)")
    return texts


def pdf_to_txt(paths: list[Path]) -> list[Path]:
    outs = []
    for p in paths:
        text = "\n\n".join(t.replace("\r\n", "\n") for t in _pdf_pages_text(p))
        out = out_file(p, "txt")
        out.write_text(text, encoding="utf-8")
        outs.append(out)
    return outs


def pdf_to_docx(paths: list[Path]) -> list[Path]:
    from docx import Document
    from docx.enum.text import WD_BREAK

    outs = []
    for p in paths:
        pages = _pdf_pages_text(p)
        d = Document()
        for k, page in enumerate(pages):
            for para in page.replace("\r\n", "\n").split("\n\n"):
                if para.strip():
                    d.add_paragraph(para.replace("\n", " ").strip())
            if k < len(pages) - 1:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        out = out_file(p, "docx")
        d.save(out)
        outs.append(out)
    return outs


def _reader(path: Path):
    from pypdf import PdfReader

    try:
        r = PdfReader(str(path))
    except Exception:
        raise Oops(f"Impossible d'ouvrir {path.name}")
    if r.is_encrypted:
        raise Oops(f"{path.name} est protégé par un mot de passe")
    return r


def pdf_compress(paths: list[Path]) -> list[Path]:
    from pypdf import PdfWriter

    outs = []
    for p in paths:
        w = PdfWriter(clone_from=_reader(p))
        for page in w.pages:
            try:
                for img in page.images:
                    pil = img.image
                    if pil.mode not in ("RGB", "L"):
                        pil = pil.convert("RGB")
                    pil.thumbnail((1600, 1600))
                    img.replace(pil, quality=60)
            except Exception:
                pass
            try:
                page.compress_content_streams()
            except Exception:
                pass
        try:
            w.compress_identical_objects(remove_identicals=True, remove_orphans=True)
        except Exception:
            pass
        out = out_file(p, "pdf", " (compressé)")
        with open(out, "wb") as f:
            w.write(f)
        if size(out) >= size(p):
            discard(out)
            continue
        outs.append(out)
    if not outs:
        raise Oops("Ce PDF est déjà optimisé")
    return outs


def pdf_merge(paths: list[Path]) -> list[Path]:
    from pypdf import PdfWriter

    w = PdfWriter()
    for p in sorted(paths, key=lambda x: x.name.lower()):
        w.append(_reader(p))
    out = unique(paths[0].parent, "Fusion", "pdf")
    with open(out, "wb") as f:
        w.write(f)
    return [out]


def pdf_split(paths: list[Path]) -> list[Path]:
    from pypdf import PdfWriter

    outs = []
    for p in paths:
        r = _reader(p)
        if len(r.pages) < 2:
            raise Oops(f"{p.name} n'a qu'une seule page")
        folder = out_folder(p, " (pages)")
        for i, page in enumerate(r.pages):
            w = PdfWriter()
            w.add_page(page)
            with open(folder / f"Page {i + 1}.pdf", "wb") as f:
                w.write(f)
        outs.append(folder)
    return outs


def pdf_rotate(paths: list[Path]) -> list[Path]:
    from pypdf import PdfWriter

    outs = []
    for p in paths:
        w = PdfWriter(clone_from=_reader(p))
        for page in w.pages:
            page.rotate(90)
        out = out_file(p, "pdf", " (pivoté)")
        with open(out, "wb") as f:
            w.write(f)
        outs.append(out)
    return outs


def pdf_strip(paths: list[Path]) -> list[Path]:
    from pypdf import PdfWriter

    outs = []
    for p in paths:
        w = PdfWriter()
        for page in _reader(p).pages:
            w.add_page(page)
        try:
            w.metadata = None
        except Exception:
            w.add_metadata({"/Producer": ""})
        out = out_file(p, "pdf", " (sans métadonnées)")
        with open(out, "wb") as f:
            w.write(f)
        outs.append(out)
    return outs


# ---------------------------------------------------------------- vidéo / audio

VIDEO_CODECS = {
    "mp4": ["-c:v", "libx264", "-preset", "medium", "-crf", "20", "-pix_fmt", "yuv420p",
            "-c:a", "aac", "-b:a", "192k", "-movflags", "+faststart"],
    "mov": ["-c:v", "libx264", "-preset", "medium", "-crf", "20", "-pix_fmt", "yuv420p",
            "-c:a", "aac", "-b:a", "192k"],
    "webm": ["-c:v", "libvpx-vp9", "-crf", "32", "-b:v", "0", "-c:a", "libopus", "-b:a", "128k"],
    "mkv": ["-c:v", "libx264", "-crf", "20", "-c:a", "aac", "-b:a", "192k"],
}

AUDIO_CODECS = {
    "mp3": ["-vn", "-c:a", "libmp3lame", "-q:a", "2"],
    "m4a": ["-vn", "-c:a", "aac", "-b:a", "256k"],
    "wav": ["-vn", "-c:a", "pcm_s16le"],
    "flac": ["-vn", "-c:a", "flac"],
    "ogg": ["-vn", "-c:a", "libvorbis", "-q:a", "5"],
}


def media_convert(paths: list[Path], ext: str) -> list[Path]:
    codecs = VIDEO_CODECS.get(ext) or AUDIO_CODECS[ext]
    outs = []
    for p in paths:
        out = out_file(p, ext)
        run_ffmpeg(["-i", str(p), *codecs], out)
        outs.append(out)
    return outs


def video_gif(paths: list[Path]) -> list[Path]:
    outs = []
    vf = "fps=12,scale='min(540,iw)':-1:flags=lanczos,split[a][b];[a]palettegen[p];[b][p]paletteuse"
    for p in paths:
        out = out_file(p, "gif")
        run_ffmpeg(["-t", "20", "-i", str(p), "-vf", vf, "-loop", "0"], out)
        outs.append(out)
    return outs


def video_compress(paths: list[Path]) -> list[Path]:
    outs = []
    for p in paths:
        out = out_file(p, "mp4", " (compressé)")
        run_ffmpeg(["-i", str(p), "-vf", "scale='min(1920,iw)':-2", "-c:v", "libx264", "-preset", "medium",
                    "-crf", "28", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "128k",
                    "-movflags", "+faststart"], out)
        if size(out) >= size(p):
            discard(out)
            continue
        outs.append(out)
    if not outs:
        raise Oops("Vidéo déjà optimisée")
    return outs


def video_mute(paths: list[Path]) -> list[Path]:
    outs = []
    for p in paths:
        ext = norm_ext(p.suffix) or "mp4"
        out = out_file(p, ext, " (sans son)")
        run_ffmpeg(["-i", str(p), "-c", "copy", "-an"], out)
        outs.append(out)
    return outs


def video_frame(paths: list[Path]) -> list[Path]:
    outs = []
    for p in paths:
        out = out_file(p, "jpg", " (image)")
        try:
            run_ffmpeg(["-ss", "1", "-i", str(p), "-frames:v", "1", "-q:v", "2"], out)
        except Oops:
            run_ffmpeg(["-i", str(p), "-frames:v", "1", "-q:v", "2"], out)
        if not out.exists():
            run_ffmpeg(["-i", str(p), "-frames:v", "1", "-q:v", "2"], out)
        outs.append(out)
    return outs


def audio_compress(paths: list[Path]) -> list[Path]:
    outs = []
    for p in paths:
        out = out_file(p, "mp3", " (compressé)")
        run_ffmpeg(["-i", str(p), "-vn", "-c:a", "libmp3lame", "-b:a", "128k"], out)
        outs.append(out)
    return outs


# ---------------------------------------------------------------- documents (Qt, fil principal)

def _load_document(path: Path):
    from PySide6.QtGui import QTextDocument

    doc = QTextDocument()
    ext = norm_ext(path.suffix)
    if ext == "docx":
        from docx import Document
        import html as _html

        parts = []
        for para in Document(str(path)).paragraphs:
            style = (para.style.name or "").lower() if para.style is not None else ""
            runs = []
            for r in para.runs:
                t = _html.escape(r.text)
                if r.bold:
                    t = f"<b>{t}</b>"
                if r.italic:
                    t = f"<i>{t}</i>"
                if r.underline:
                    t = f"<u>{t}</u>"
                runs.append(t)
            inner = "".join(runs) or "&nbsp;"
            if style.startswith("heading") or style.startswith("titre"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                parts.append(f"<h{level}>{inner}</h{level}>")
            elif style == "title":
                parts.append(f"<h1>{inner}</h1>")
            else:
                parts.append(f"<p>{inner}</p>")
        doc.setHtml("\n".join(parts))
        return doc
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    if ext == "md":
        doc.setMarkdown(text)
    elif ext == "html":
        doc.setHtml(text)
    else:
        doc.setPlainText(text)
    return doc


def document_convert(paths: list[Path], ext: str) -> list[Path]:
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QFont, QPageLayout, QPageSize, QPdfWriter

    outs = []
    for p in paths:
        doc = _load_document(p)
        out = out_file(p, ext)
        if ext == "pdf":
            doc.setDefaultFont(QFont("Arial", 11))
            writer = QPdfWriter(str(out))
            writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
            writer.setPageMargins(QMarginsF(20, 20, 20, 20), QPageLayout.Unit.Millimeter)
            writer.setResolution(300)
            doc.print_(writer)
            del writer
        elif ext == "txt":
            out.write_text(doc.toPlainText(), encoding="utf-8")
        elif ext == "md":
            out.write_text(doc.toMarkdown(), encoding="utf-8")
        elif ext == "html":
            out.write_text(doc.toHtml(), encoding="utf-8")
        elif ext == "docx":
            from docx import Document

            d = Document()
            for para in doc.toPlainText().split("\n"):
                d.add_paragraph(para)
            d.save(out)
        else:
            raise Oops(f"Format {ext.upper()} non pris en charge")
        outs.append(out)
    return outs


# ---------------------------------------------------------------- archives

def make_zip(paths: list[Path]) -> list[Path]:
    if len(paths) == 1:
        out = unique(paths[0].parent, paths[0].name if paths[0].is_dir() else paths[0].stem, "zip")
    else:
        out = unique(paths[0].parent, "Archive", "zip")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            if p.is_dir():
                for root, _, files in os.walk(p):
                    for f in files:
                        full = Path(root) / f
                        z.write(full, full.relative_to(p.parent))
            else:
                z.write(p, p.name)
    return [out]
