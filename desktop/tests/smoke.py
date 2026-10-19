"""Test automatique : essaie chaque action sur des fichiers d'exemple et dessine la roue.

Lancé par GitHub Actions (QT_QPA_PLATFORM=offscreen). Sort en erreur si une action échoue.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from PIL import Image, ImageDraw  # noqa: E402
from PySide6.QtCore import QPointF, Qt  # noqa: E402
from PySide6.QtGui import QColor, QPainter, QPixmap  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from aikoroue import catalog, converters as c, wheel  # noqa: E402

# Erreurs « normales » (le fichier ne s'y prête pas) : pas un échec du test
EXPECTED = ("déjà optimis", "qu'une seule page")


def make_samples(d: Path) -> dict[str, list[Path]]:
    im = Image.new("RGB", (1600, 1000), (240, 140, 60))
    ImageDraw.Draw(im).ellipse([500, 200, 1100, 800], fill=(40, 90, 200))
    im.save(d / "photo.jpg", quality=95)
    logo = Image.new("RGBA", (800, 800), (0, 0, 0, 0))
    ImageDraw.Draw(logo).rectangle([200, 200, 600, 600], fill=(255, 80, 80, 255))
    logo.save(d / "logo.png")

    (d / "notes.txt").write_text("Bonjour AikoRoue !\n\nCeci est un test de conversion.\n", encoding="utf-8")
    (d / "notes.md").write_text("# Titre\n\nUn paragraphe **gras** et *italique*.\n\n- un\n- deux\n", encoding="utf-8")
    from docx import Document

    doc = Document()
    doc.add_heading("Document de test", 1)
    doc.add_paragraph("Premier paragraphe.")
    doc.add_paragraph("Deuxième paragraphe.")
    doc.save(d / "lettre.docx")

    # PDF de 2 pages avec du vrai texte (via Qt)
    long_md = "\n\n".join(f"## Section {i}\n\n" + "Du texte de test. " * 60 for i in range(1, 8))
    (d / "long.md").write_text(long_md, encoding="utf-8")
    pdf = c.document_convert([d / "long.md"], "pdf")[0]
    pdf.rename(d / "rapport.pdf")
    shutil.copy(d / "rapport.pdf", d / "annexe.pdf")
    (d / "long.md").unlink()

    ff = c.ffmpeg_exe()
    subprocess.run([ff, "-hide_banner", "-loglevel", "error", "-y",
                    "-f", "lavfi", "-i", "testsrc=duration=3:size=640x360:rate=25",
                    "-f", "lavfi", "-i", "sine=frequency=440:duration=3",
                    "-shortest", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac",
                    str(d / "clip.mp4")], check=True)
    subprocess.run([ff, "-hide_banner", "-loglevel", "error", "-y",
                    "-f", "lavfi", "-i", "sine=frequency=330:duration=4", str(d / "son.wav")], check=True)

    return {
        "photo": [d / "photo.jpg"],
        "logo": [d / "logo.png"],
        "pdf": [d / "rapport.pdf"],
        "2 pdf": [d / "rapport.pdf", d / "annexe.pdf"],
        "vidéo": [d / "clip.mp4"],
        "audio": [d / "son.wav"],
        "texte": [d / "notes.txt"],
        "markdown": [d / "notes.md"],
        "docx": [d / "lettre.docx"],
    }


def render_wheels(out_dir: Path, samples: dict[str, list[Path]]):
    out_dir.mkdir(parents=True, exist_ok=True)
    for dark in (True, False):
        wheel.is_dark = lambda dark=dark: dark
        for name, hover in (("image", (0, 1)), ("pdf", (1, 2)), ("video", (0, 0))):
            paths = {"image": samples["photo"], "pdf": samples["pdf"], "video": samples["vidéo"]}[name]
            w = wheel.WheelWindow()
            w._configure(paths)
            w.hovered = hover
            pm = w.grab()
            canvas = QPixmap(pm.size())
            canvas.fill(QColor("#2B2F36") if dark else QColor("#D9DEE5"))
            p = QPainter(canvas)
            p.drawPixmap(0, 0, pm)
            p.end()
            canvas.save(str(out_dir / f"roue-{name}-{'sombre' if dark else 'clair'}.png"))
        t = wheel.Toast()
        t.show_message("✓ photo.png", False, QPointF(300, 300).toPoint())
        t.grab().save(str(out_dir / f"bulle-{'sombre' if dark else 'clair'}.png"))
        t.hide()


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    app = QApplication(sys.argv)
    failures = []
    shots = ROOT / "screenshots"
    with tempfile.TemporaryDirectory() as tmp:
        d = Path(tmp)
        samples = make_samples(d)
        render_wheels(shots, samples)
        for label, paths in samples.items():
            acts = catalog.actions_for(paths)
            print(f"\n== {label} ({', '.join(p.name for p in paths)}) : {[a.title for a in acts]}")
            for a in acts:
                try:
                    outs = a.run(paths)
                    missing = [o for o in outs if not Path(o).exists()]
                    if missing:
                        raise RuntimeError(f"fichier absent : {missing}")
                    print(f"  ok   {a.title:12} → {', '.join(Path(o).name for o in outs)}")
                except Exception as e:  # noqa: BLE001
                    msg = str(e)
                    if any(x in msg for x in EXPECTED):
                        print(f"  --   {a.title:12} ({msg})")
                    else:
                        print(f"  FAIL {a.title:12} {type(e).__name__}: {msg}")
                        failures.append(f"{label}/{a.title}: {type(e).__name__}: {msg}")
                        if os.environ.get("GITHUB_ACTIONS"):
                            print(f"::error title={label} / {a.title}::{type(e).__name__}: {msg}".replace("\n", " "))
    print("\nCaptures :", ", ".join(sorted(p.name for p in shots.glob("*.png"))))
    if failures:
        print("\nÉCHECS :\n" + "\n".join(failures))
        return 1
    print("\nTout est OK")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001
        import traceback

        tb = traceback.format_exc()
        print(tb)
        if os.environ.get("GITHUB_ACTIONS"):
            print("::error title=Erreur fatale::" + " | ".join(tb.strip().splitlines()[-6:]))
        sys.exit(1)
